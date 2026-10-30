"""Builds the kit to send to the team running the manual-entry comparison.

  python3 build_session_kit.py

Writes to session_kit/
  Verity_Manual_Comparison_Checklist.docx   step-by-step checklist for the facilitator
  Verity_Session_Pack.zip                   order messages, reference sheets, entry template, order sets and schedule.
                                            The answer key, the SQL and the scripts are deliberately left out.
"""
import csv
import zipfile
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt

HERE = Path(__file__).parent
OUT = HERE / 'session_kit'
OUT.mkdir(exist_ok=True)
sets = list(csv.DictReader(open(HERE / 'sample_80.csv')))
session_date = open(HERE / 'messages_all.txt', encoding='utf-8').read().split('Date: ', 1)[1].split(' 09:')[0] if 'Date: ' in open(HERE / 'messages_all.txt', encoding='utf-8').read() else 'the date printed on the messages'

with open(OUT / 'order_sets.csv', 'w', newline='') as f:
    w = csv.writer(f)
    w.writerow(['reference', 'message_format', 'order_set', 'message_file'])
    for r in sets:
        w.writerow([r['reference'], r['message_format'], r['order_set'], r['message_file']])

with zipfile.ZipFile(OUT / 'Verity_Session_Pack.zip', 'w', zipfile.ZIP_DEFLATED) as z:
    for p in sorted((HERE / 'messages').iterdir()):
        z.write(p, f'messages/{p.name}')
    z.write(HERE / 'messages_all.txt', 'messages_all_printable.txt')
    for p in sorted((HERE / 'reference').iterdir()):
        z.write(p, f'reference/{p.name}')
    z.write(HERE / 'entry_template.xlsx', 'entry_template.xlsx')
    z.write(HERE / 'counterbalancing_schedule.csv', 'participant_schedule.csv')
    z.write(OUT / 'order_sets.csv', 'order_sets.csv')
    z.write(HERE.parent / '06_questionnaire' / 'consent_text.docx', 'consent_text.docx')

doc = Document()
st = doc.styles['Normal']
st.font.name, st.font.size = 'Calibri', Pt(11)
st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Calibri')


def h(t, lvl=1):
    doc.add_heading(t, level=lvl)


def p(t, bold=False):
    para = doc.add_paragraph()
    r = para.add_run(t)
    r.bold = bold
    return para


def box(items):
    for t in items:
        doc.add_paragraph('☐  ' + t)


title = doc.add_heading('Verity manual-entry comparison: session checklist', 0)
p('This session measures how many order errors get through, and how long orders take, when the same orders are handled by hand and '
  'through Verity. The results go into Chapter Four for Objectives 3 and 4, so please follow the steps exactly and record what '
  'actually happens. Do not fill anything in afterwards from memory.')
p(f'The order messages in the pack are dated {session_date}. If the session runs on another day, tell the project team first so '
  'the pack can be regenerated with the right dates.', bold=True)

h('1. Before the day')
box(['Agree the session date with the project team. They run the database setup on the morning of the session, which creates the '
     'customer logins used in the Verity part and switches automatic approval off. Without this step the Verity part cannot be scored.',
     'Recruit the participants. Each one is a real person who normally types or checks orders. Aim for at least 4, ideally 8 to 12.',
     'Give each participant an ID from P01 upwards. Use IDs only on every sheet, never names.',
     'Print or share the files in Verity_Session_Pack.zip, one copy per participant. Do not share anything else from the project folder.',
     'Prepare one computer per participant with a browser and Excel, and one extra computer for the Ops Desk approver.',
     'Choose one person to act as the Ops Desk approver for the Verity part. This person does not type orders.',
     'Have a clock or phone visible to every participant, showing hours, minutes and seconds.'])

h('2. At the start of the session')
box(['Fill in the project contact e-mail in consent_text.docx, then read it to each participant and record their consent. No consent, no participation.',
     'Look up each participant in participant_schedule.csv. It says which condition they do first, Manual or Verity, and which order '
     'set, 1 or 2. Each participant does one set by hand and the other set in Verity.',
     'Use order_sets.csv to hand each participant the messages for their set, in reference order (B-001, B-002 and so on).',
     'Give every participant the three reference files: price_list.xlsx, customer_register.xlsx and orders_already_received.csv.',
     'Explain the rules below and answer questions about the process only. Never say whether an order is right or wrong.'])

h('3. Manual part')
p('Each participant works alone, the way orders are handled today, with only their own visual check.')
box(['Open entry_template.xlsx. Use one row per order.',
     'For each message, type the Start time as HH:MM:SS the moment the message is opened.',
     'Fill in Participant ID, Condition = Manual, Order reference, Customer, Delivery address, Delivery date and the item code and '
     'quantity of each line.',
     'If anything looks wrong, set Problem flagged to Yes and describe it. Still type the order exactly as it was sent. Do not correct it.',
     'Type the End time as HH:MM:SS when the row is finished. The Seconds column fills itself.',
     'No colleague checks the work and nobody helps. Move straight on to the next message.',
     'At the end, save the file as manual_P01.xlsx, manual_P02.xlsx and so on.'])

h('4. Verity part')
p('Each participant enters their other set of orders as the customer who sent each message.')
box(['Sign in as the customer named in the message, using the e-mail from customer_register.xlsx and the password given by the project '
     'team on the day. Sign out and sign in again when the next message comes from a different customer.',
     'Create a new order. Add each line, choose the saved address and the delivery date exactly as the message gives them.',
     'In the Notes field type the order reference and the participant ID, for example B-017 P03. This is how the results are matched, '
     'so it must be on every order.',
     'Submit. If Verity blocks the order or asks to confirm a possible repeat, read the message and move on. Do not fix the order and do '
     'not resubmit it.',
     'Participants do not time anything in this part. Verity records the times itself.',
     'The Ops Desk approver signs in with the Ops account and works through the approval queue in order, approving or rejecting each '
     'order. When rejecting, write the reason in the comment.'])

h('5. Things that must not happen')
box(['Nobody sees an answer key or is told which orders contain mistakes.',
     'Nobody corrects an order before typing it, in either part.',
     'Nobody types times in afterwards. If a time was missed, leave it blank and note it.',
     'No real customer data is used. Only the messages in the pack.',
     'Names of participants do not appear in any file that is sent back.'])

h('6. After the session')
box(['Collect every manual_Pxx.xlsx file.',
     'Write down anything that went wrong, such as a computer problem, a missed time or a participant who stopped early.',
     'Tell the project team the session has finished. They extract the Verity results from the database and then remove the session '
     'logins. Nobody else should delete anything.',
     'Send the project team the manual files, the note of problems, and the number of participants who took part.'])

h('7. What to send back')
box(['All manual_Pxx.xlsx files.',
     'A short note with the date, the start and end time of the session, the number of participants and any problems.',
     'Confirmation that consent was recorded for every participant. Keep the consent records yourself and do not send names.'])

for para in doc.paragraphs:
    if para.style.name == 'Normal':
        para.paragraph_format.space_after = Pt(4)
doc.save(OUT / 'Verity_Manual_Comparison_Checklist.docx')
print('written', sorted(x.name for x in OUT.iterdir()), 'session date', session_date)
