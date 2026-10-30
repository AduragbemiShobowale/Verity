"""Scores the manual-entry sheets against the answer key and compares them with Verity's results on the same orders
from simulation replicate 1 (run 5), from which the pack orders were drawn. Standard library plus openpyxl.

  python3 score_manual_vs_simulation.py session_data/raw/*.xlsx

Uses the definitions fixed in analyze_baseline.py before any data were collected, with two stated adaptations.
  1. The Verity side is the simulation outcome for the source order (answer key column source_order_no_run5), not a
     person entering the order in Verity. In Verity the customer enters the order against the live catalogue, so no
     staff retyping takes place and no typing error is counted on the Verity side. An error reaches the record in
     Verity when a planted error was not flagged (outcome pass).
  2. A delivery date stored with day and month swapped counts as typed correctly when the swap matches the message.
     Google Sheets with a US locale reads 03/11/2026 typed as DD/MM as 11 March, so the swap is the sheet's, not the
     participant's.
Times: manual seconds are End time minus Start time as typed. Verity's human entry time was not measured, so no time
saving is computed. Writes session_data/manual_vs_simulation_results.json.
"""
import csv
import datetime as dt
import json
import math
import statistics as st
from collections import Counter
import sys
from pathlib import Path
import openpyxl

HERE = Path(__file__).parent
Z = 1.959964
key = {r['reference']: r for r in csv.DictReader(open(HERE / 'facilitator_only' / 'ANSWER_KEY.csv'))}
sim = {r['order_no']: r for r in csv.DictReader(open(HERE.parent / '03_simulation' / 'replicate1_orders.csv'))}


def wilson(k, n):
    p = k / n
    d = 1 + Z * Z / n
    c = (p + Z * Z / (2 * n)) / d
    h = Z * math.sqrt(p * (1 - p) / n + Z * Z / (4 * n * n)) / d
    return [max(0.0, round(100 * (c - h), 1)), round(100 * (c + h), 1)]


def norm(s):
    return ' '.join(str(s or '').lower().replace(',', ' ').split())


def lines_sent(s):
    out = {}
    for part in (s or '').split(';'):
        if ' x ' in part:
            sku, q = part.strip().rsplit(' x ', 1)
            out[sku.strip().upper()] = float(q)
    return out


def date_ok(v, sent):
    if not sent:
        return v in (None, '')
    want = dt.date.fromisoformat(sent)
    if isinstance(v, dt.datetime):
        v = v.date()
    if isinstance(v, dt.date):
        return v == want or (v.day <= 12 and dt.date(v.year, v.day, v.month) == want)
    try:
        return dt.datetime.strptime(str(v).strip(), '%d/%m/%Y').date() == want
    except ValueError:
        return False


def secs(a, b):
    if not isinstance(a, dt.time) or not isinstance(b, dt.time):
        return None
    return (dt.datetime.combine(dt.date.min, b) - dt.datetime.combine(dt.date.min, a)).total_seconds()


rows = []
for f in sys.argv[1:]:
    ws = openpyxl.load_workbook(f).active
    head = [c.value for c in ws[1]]
    for vals in ws.iter_rows(min_row=2, values_only=True):
        r = dict(zip(head, vals))
        ref = (r.get('Order reference') or '').strip()
        if not ref or r.get('Customer') is None:
            continue
        k = key[ref]
        lines = {}
        for i in range(1, 5):
            sku = (r.get(f'Line {i} item code') or '').strip().upper()
            if sku:
                lines[sku] = float(r.get(f'Line {i} quantity'))
        typed = {'customer': norm(r.get('Customer')) != norm(k['customer']),
                 'address': norm(r.get('Delivery address')) != norm(k['address_as_sent']),
                 'date': not date_ok(r.get('Delivery date (DD/MM/YYYY)'), k['delivery_date_as_sent']),
                 'lines': lines != lines_sent(k['lines_as_sent'])}
        flagged = str(r.get('Problem flagged (Yes/No)') or '').strip().lower() == 'yes'
        planted = k['planted_error'] != 'clean'
        s = sim[k['source_order_no_run5']]
        v_flag = s['flagged'] == 'Yes'
        abbrev_only = typed['address'] and norm(r.get('Delivery address')).replace(' rd ', ' road ') == norm(k['address_as_sent']) and not any(v for x, v in typed.items() if x != 'address')
        rows.append({'ref': ref, 'abbreviation_only': bool(abbrev_only and not (planted and not flagged)), 'participant': r.get('Participant ID'), 'planted': k['planted_error'], 'has_rule': k['has_rule'],
                     'manual_flagged': flagged, 'manual_description': r.get('Problem description'),
                     'typing_errors': [x for x, bad in typed.items() if bad],
                     'manual_error_reached': (planted and not flagged) or any(typed.values()),
                     'manual_planted_missed': planted and not flagged, 'manual_false_flag': flagged and not planted,
                     'manual_seconds': secs(r.get('Start time (HH:MM:SS)'), r.get('End time (HH:MM:SS)')),
                     'verity_order': s['order_no'], 'verity_outcome': s['outcome'], 'verity_flagged': v_flag,
                     'verity_error_reached': planted and not v_flag, 'verity_false_flag': v_flag and not planted,
                     'verity_validation_ms': float(s['validation_ms'])})


def rate(k, n):
    return {'k': k, 'n': n, 'pct': round(100 * k / n, 1), 'ci95': wilson(k, n)}


n = len(rows)
err = [r for r in rows if r['planted'] != 'clean']
clean = [r for r in rows if r['planted'] == 'clean']
mk = sum(r['manual_error_reached'] for r in rows)
vk = sum(r['verity_error_reached'] for r in rows)
ms = [r['manual_seconds'] for r in rows if r['manual_seconds'] is not None]
out = {
    'orders': n, 'participants': sorted({r['participant'] for r in rows}),
    'planted_errors': len(err), 'clean_orders': len(clean),
    'manual_error_rate': rate(mk, n), 'verity_error_rate': rate(vk, n),
    'error_reduction_pct': round(100 * (1 - (vk / n) / (mk / n)), 1) if mk else None,
    'manual_planted_errors_caught': rate(sum(r['manual_flagged'] for r in err), len(err)),
    'verity_planted_errors_flagged': rate(sum(r['verity_flagged'] for r in err), len(err)),
    'manual_caught_rule_targeted': rate(sum(r['manual_flagged'] for r in err if r['has_rule'] == 'Yes'), sum(r['has_rule'] == 'Yes' for r in err)),
    'verity_flagged_rule_targeted': rate(sum(r['verity_flagged'] for r in err if r['has_rule'] == 'Yes'), sum(r['has_rule'] == 'Yes' for r in err)),
    'verity_errors_passed_by_type': dict(sorted(Counter(r['planted'] for r in err if not r['verity_flagged']).items())),
    'manual_rate_excluding_abbreviation_only': None,
    'manual_caught_no_rule': rate(sum(r['manual_flagged'] for r in err if r['has_rule'] == 'No'), sum(r['has_rule'] == 'No' for r in err)),
    'verity_flagged_no_rule': rate(sum(r['verity_flagged'] for r in err if r['has_rule'] == 'No'), sum(r['has_rule'] == 'No' for r in err)),
    'manual_orders_with_typing_error': rate(sum(bool(r['typing_errors']) for r in rows), n),
    'typing_errors_by_field': {f: sum(f in r['typing_errors'] for r in rows) for f in ('customer', 'address', 'date', 'lines')},
    'manual_false_flags_on_clean': rate(sum(r['manual_false_flag'] for r in clean), len(clean)),
    'verity_false_flags_on_clean': rate(sum(r['verity_false_flag'] for r in clean), len(clean)),
    'manual_seconds_flagged_median': st.median(r['manual_seconds'] for r in rows if r['manual_flagged']),
    'manual_seconds_unflagged_median': st.median(r['manual_seconds'] for r in rows if not r['manual_flagged']),
    'manual_seconds': {'orders': len(ms), 'mean': round(st.mean(ms), 1), 'median': round(st.median(ms), 1), 'sd_sample': round(st.stdev(ms), 1),
                       'min': min(ms), 'max': max(ms), 'total_minutes': round(sum(ms) / 60, 1)},
    'verity_validation_ms_same_orders': {'mean': round(st.mean(r['verity_validation_ms'] for r in rows), 2),
                                         'median': round(st.median(r['verity_validation_ms'] for r in rows), 2)},
    'indicative_only': True,
    'rows': rows,
}
# Sensitivity: an order whose only counted error is an abbreviated address (Rd for Road) would not misdirect a delivery
mk2 = mk - sum(r['abbreviation_only'] for r in rows)
out['manual_rate_excluding_abbreviation_only'] = rate(mk2, n)
out['error_reduction_excluding_abbreviation_only_pct'] = round(100 * (1 - (vk / n) / (mk2 / n)), 1)
out['objective_3_error_reduction_status'] = 'Met' if out['error_reduction_pct'] is not None and out['error_reduction_pct'] >= 60 else 'Not met'
(HERE / 'session_data' / 'manual_vs_simulation_results.json').write_text(json.dumps(out, indent=1, default=str))
print(json.dumps({k: v for k, v in out.items() if k != 'rows'}, indent=1))
for r in rows:
    print(r['ref'], r['participant'], r['planted'], 'M-flag' if r['manual_flagged'] else '-', r['typing_errors'], 'V-flag' if r['verity_flagged'] else '-', r['manual_seconds'])
