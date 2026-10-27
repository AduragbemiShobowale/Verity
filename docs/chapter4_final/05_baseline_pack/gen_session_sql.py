"""Writes setup_baseline_session.sql and cleanup_baseline_session.sql from the pack's reference files.
Neither script has been run against the live database. They are for the facilitator, on the day of the session."""
import csv
from pathlib import Path
from openpyxl import load_workbook

HERE = Path(__file__).parent
reg = list(load_workbook(HERE / 'reference' / 'customer_register.xlsx').active.iter_rows(min_row=2, values_only=True))
log = list(csv.DictReader(open(HERE / 'reference' / 'orders_already_received.csv')))
q = lambda s: "'" + str(s).replace("'", "''") + "'"
STATUS = {'Active': 'active', 'On hold': 'on_hold', 'Blocked': 'blocked'}
cust_rows = ',\n    '.join(f"({q(c)}, {q(ct)}, {q(e)}, {q(ph)}, {q(a1)}, {q(a2)}, {q(rg)}, {q(STATUS[st])})" for c, ct, e, ph, a1, a2, rg, st in reg)
log_rows = ',\n    '.join(f"({q(r['customer'])}, {q(r['lines'])})" for r in log)
setup = f"""-- Baseline session setup for the Verity condition. NOT RUN. Run once, as the database owner, on the morning of the session.
-- Creates one dummy login per fictional customer in the pack (password below), the two saved addresses, the account
-- status shown in the customer register, four earlier delivered orders per customer (as in the simulation, so the
-- unusual-quantity rule has a history), and the orders listed in orders_already_received.csv as submitted this morning.
-- It also turns automatic approval off and marks SAND-T discontinued for the session. cleanup_baseline_session.sql undoes it.
-- All addresses use the reserved .example domain. No real customer is touched.
do $$
declare
  r record; uid uuid; cid uuid; addr uuid; o uuid; k int; it text; items jsonb;
  pw text := 'Baseline#2026';
begin
  update settings set value = 'false' where key = 'auto_approve_enabled';
  update products set active = false where sku = 'SAND-T';
  for r in select * from (values
    {cust_rows}
  ) v(company, contact, email, phone, addr1, addr2, region, status) loop
    uid := _seed_user(r.email, r.contact, 'customer', pw);
    insert into customers (name, reg_no, contact_name, email, phone, address, industry, expected_volume, status, admin_note)
    values (r.company, 'RC-BASE-' || substr(md5(r.email), 1, 6), r.contact, r.email, r.phone, r.addr1, 'Construction', 20, 'active', 'baseline session dummy')
    returning id into cid;
    insert into customer_users values (uid, cid);
    insert into addresses (customer_id, label, line, region_id) values
      (cid, 'Head office', r.addr1, (select id from regions where name = r.region)),
      (cid, 'Site', r.addr2, (select id from regions where name = r.region));
    perform _as(r.email);
    addr := (select id from addresses where customer_id = cid and label = 'Head office');
    for k in 1..4 loop
      o := save_draft(null, jsonb_build_object('address_id', addr, 'requested_date', current_date + 3, 'notes', 'baseline history',
             'items', jsonb_build_array(jsonb_build_object('product_id', (select id from products where sku = 'CEM-50'), 'qty', 20 + k),
                                        jsonb_build_object('product_id', (select id from products where sku = 'ROD-12'), 'qty', 10 + k))));
      perform submit_order(o, true);
      -- history only: move it out of the approval queue and the 24-hour duplicate window
      update orders set status = 'delivered', submitted_at = now() - interval '14 days' where id = o;
    end loop;
    update customers set status = r.status::customer_status where id = cid;
  end loop;
  -- orders already received this morning (the reference log), so repeats can be recognised
  for r in select * from (values
    {log_rows}
  ) v(company, lines) loop
    cid := (select id from customers where name = r.company and admin_note = 'baseline session dummy');
    perform _as((select email from customers where id = cid));
    items := '[]';
    foreach it in array string_to_array(r.lines, '; ') loop
      items := items || jsonb_build_array(jsonb_build_object('product_id', (select id from products where sku = split_part(it, ' x ', 1)),
                                                             'qty', split_part(it, ' x ', 2)::numeric));
    end loop;
    o := save_draft(null, jsonb_build_object('address_id', (select id from addresses where customer_id = cid and label = 'Head office'),
           'requested_date', current_date + 3, 'notes', 'baseline earlier order', 'items', items));
    perform submit_order(o, true);
  end loop;
end $$;
"""
cleanup = """-- Undo setup_baseline_session.sql after the session and after extract_baseline.sql has been saved. NOT RUN.
-- Removes only rows that belong to the dummy customers (admin_note = 'baseline session dummy').
begin;
update settings set value = 'true' where key = 'auto_approve_enabled';
update products set active = true where sku = 'SAND-T';
create temp table b_cust as select id from customers where admin_note = 'baseline session dummy';
create temp table b_ord as select id from orders where customer_id in (select id from b_cust);
create temp table b_user as select user_id from customer_users where customer_id in (select id from b_cust);
-- stock reserved by approved baseline orders is released before the orders are removed
update products p set reserved = greatest(p.reserved - x.q, 0)
  from (select oi.product_id, sum(oi.qty_ordered) q from order_items oi join orders o on o.id = oi.order_id
        where o.id in (select id from b_ord) and o.status in ('awaiting_payment','payment_verification','paid','picking','packed') group by 1) x
  where p.id = x.product_id;
delete from notifications where order_id in (select id from b_ord);
delete from validation_results where order_id in (select id from b_ord);
delete from approvals where order_id in (select id from b_ord);
delete from invoices where order_id in (select id from b_ord);
delete from order_events where order_id in (select id from b_ord);
delete from order_items where order_id in (select id from b_ord);
delete from orders where id in (select id from b_ord);
delete from addresses where customer_id in (select id from b_cust);
delete from customer_users where customer_id in (select id from b_cust);
delete from customers where id in (select id from b_cust);
delete from auth.users where id in (select user_id from b_user);
commit;
"""
(HERE / 'setup_baseline_session.sql').write_text(setup)
(HERE / 'cleanup_baseline_session.sql').write_text(cleanup)
print('written', len(reg), 'customers', len(log), 'earlier orders')
